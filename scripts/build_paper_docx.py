"""Create the editable research report with original Python-generated PNGs."""

from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build():
    root = Path(__file__).resolve().parents[1]
    output = root / "docs" / "copart-research-report.docx"
    figures = root / "docs" / "figures"
    images = [
        figures / "04-editorial-weather-distribution.png",
        figures / "02-weather-seasonality.png",
    ]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)
    section.footer_distance = Inches(0.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.02
    for name in ["Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string("253F52")
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)
    doc.styles["Title"].font.name = "Cambria"
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].paragraph_format.space_after = Pt(5)

    def para(text, size=None, bold=False, color=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def heading(text):
        doc.add_heading(text, level=1)

    doc.add_paragraph("Copart investment research report", "Title")
    para("CPRT  |  September 30, 2026  |  Research draft", 10, True)
    para("Recommendation: [Steven]   Horizon: [3–12 months]   Target: [after testing]", 10)
    para(
        "AI-assisted research draft. The team must independently author final submission text under the competition rules.",
        8,
        color="7F4B35",
    )

    heading("Investment thesis")
    para(
        "[Steven: State LONG or SHORT, the expectation the market has wrong, the evidence for a different outcome, and the catalyst that will reveal it within the investment horizon.]",
        color="895824",
    )
    para(
        "The research question is whether Copart’s operating slowdown reflects temporary conditions or a lasting change in growth and profitability. The separate Excel submission translates service revenue, vehicle sales, operating costs and reinvestment into cash flow and value. The recommendation remains open until the thesis assumptions are tested."
    )

    heading("Operating evidence and competition")
    para(
        "The supplied workbook records FY2026 revenue of $4,666.2 million, up 0.4%, and operating income of $1,652.6 million versus $1,696.7 million in FY2025. Operating margin falls from 36.5% to 35.4%. Service revenue is approximately flat at $3,969.5 million; purchased-vehicle revenue rises 2.7% to $696.7 million. FY2026 is labeled unaudited. These figures focus the thesis on growth, cost absorption and reinvestment. [1]"
    )
    para(
        "[Steven: Add cited competitive and customer-retention evidence, management assessment, the strongest alternative explanation, and why the evidence favors the thesis.]",
        color="895824",
    )

    heading("Valuation and assumption testing")
    para(
        "The workbook forecasts FY2027–FY2031 free cash flow to the firm, discounts the explicit cash flows and terminal value, and bridges enterprise value to equity value. Its provisional inputs include 9.0% WACC, 2.5% terminal growth and approximately 0.02% annual service-revenue growth. The saved standalone value is $23.30 per share, 14.4% below the workbook’s $27.22 reference price labeled September 28, 2026. This is not a validated target or a short recommendation. [2]"
    )
    para(
        "Terminal value accounts for 73.1% of enterprise value. The cash/debt/share bridge remains provisional and the acquisition-adjusted value is blank. The paper should report tested bear/base/bull assumptions and distinguish present DCF value from the eventual price target. Excel will be submitted separately. [2]"
    )
    para(
        "[After testing: insert scenario values and returns, changed growth/cost assumptions, WACC sensitivity, verified reference price, and the rationale for the selected investment horizon.]",
        color="895824",
    )

    heading("Catalysts and risks")
    para(
        "[Steven: Add two dated catalysts, the operating metric each should reveal, and a clear thesis-invalidating observation.]",
        color="895824",
    )
    para(
        "Test persistent service-revenue weakness, higher costs and greater capital needs. The workbook excludes the pending ACV transaction from standalone value; any transaction case must incorporate both consideration paid and supported value received, including integration costs. If recommending a short, address borrow and squeeze risk. [2]"
    )
    para(
        "Sources: [1] CopartDCFValuationModeL.xlsx, Historical Data D6:E13; Forecast E37:E38. [2] Same workbook, Assumptions F6:J35; Forecast F43:J47; Valuation B5:B11, B20:B32, B37:B53. Values are saved workbook results, not a fresh Excel recalculation or independently verified market quote.",
        8,
    )

    doc.add_page_break()
    heading("Weather evidence and the predictive question")
    para(
        "NOAA records show different patterns by event type. These exhibits provide context; neither establishes an effect on Copart’s vehicle volumes or earnings. The four annual workbook observations cannot support the planned predictive test without longer financial history.",
        9,
    )
    for image, width, caption, alt in [
        (
            images[0],
            5.65,
            "Figure 1. FY2025 quarters against historical same-quarter medians. Full-sample comparison, not a point-in-time backtest.",
            "Historical weather distributions with FY2025 highlighted. All four flash-flood quarters exceed their median; three hail and flood quarters are below.",
        ),
        (
            images[1],
            6.65,
            "Figure 2. Seasonal medians and middle 50% across years. Seasonality must be controlled before attributing predictive value to weather.",
            "Monthly hail, flash-flood and flood record seasonality across 2010 to 2025, with median curves and interquartile bands.",
        ),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        shape = p.add_run().add_picture(str(image), width=Inches(width))
        shape._inline.docPr.set("descr", alt)
        para(caption, 8)
    para(
        "[After validation: add forecast target, sample/test periods, baseline error and weather-model error. No validated Copart weather forecast is available yet.]",
        8,
        color="895824",
    )
    para(
        "Source: NOAA Storm Events revised 2010–2025 snapshots, U.S. states/DC; ncei.noaa.gov/pub/data/swdi/stormevents/csvfiles/. Figures are the original Python-generated PNGs. Counts are event records, not unique storms or insured losses. Historical publication vintages are not reconstructed.",
        8,
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run()
    r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r._r.addnext(fld)
    doc.core_properties.title = "Copart investment research report"
    doc.core_properties.subject = "Research draft with separate Excel model references"
    doc.core_properties.author = ""
    doc.save(output)
    with ZipFile(output) as archive:
        embedded = [
            sha256(archive.read(n)).hexdigest()
            for n in archive.namelist()
            if n.startswith("word/media/")
        ]
    assert sorted(embedded) == sorted(sha256(p.read_bytes()).hexdigest() for p in images)
    print(f"Created {output}; both embedded image hashes match the original Python PNGs.")


if __name__ == "__main__":
    build()
